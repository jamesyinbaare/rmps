"""High-throughput score import pipeline: parse → lookup maps → classify → staging apply."""

from __future__ import annotations

import csv
import hashlib
import io
import logging
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Literal

import pandas as pd
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    Candidate,
    DataExtractionMethod,
    Exam,
    ExamRegistration,
    ExamSubject,
    School,
    Subject,
    SubjectRegistration,
    SubjectScore,
)
from app.services.subject_upload import SubjectUploadParseError
from app.utils.score_utils import (
    calculate_total_score,
    is_absent,
    parse_score_value,
    validate_score_range,
)

logger = logging.getLogger(__name__)

PaperTestType = Literal[1, 2]

LARGE_IMPORT_ROW_THRESHOLD = 5000
MAX_STORED_IMPORT_ERRORS = 500
IMPORT_APPLY_BATCH_SIZE = 2000
IMPORT_PROGRESS_EVERY = 5000
IMPORT_STAGING_INSERT_BATCH = 5000
IN_CHUNK_SIZE = 10_000
SKIP_SCORE_TOKENS = frozenset({"N/A", "NA"})
METHOD = DataExtractionMethod.MANUAL_ENTRY_PHYSICAL

REQUIRED_COLUMNS = ("index_number", "subject_code", "score")
OPTIONAL_COLUMNS = ("subject_name",)

COLUMN_ALIASES: dict[str, str] = {
    "index_number": "index_number",
    "index": "index_number",
    "indexnumber": "index_number",
    "candidate_index": "index_number",
    "candidate_index_number": "index_number",
    "subject_code": "subject_code",
    "subjectcode": "subject_code",
    "code": "subject_code",
    "original_code": "subject_code",
    "score": "score",
    "raw_score": "score",
    "mark": "score",
    "marks": "score",
    "subject_name": "subject_name",
    "subjectname": "subject_name",
    "name": "subject_name",
}


@dataclass
class ScoreImportRowError:
    row: int
    index_number: str | None = None
    subject_code: str | None = None
    message: str = ""

    def as_dict(self) -> dict[str, str]:
        out: dict[str, str] = {"row": str(self.row), "message": self.message}
        if self.index_number is not None:
            out["index_number"] = self.index_number
        if self.subject_code is not None:
            out["subject_code"] = self.subject_code
        return out


@dataclass
class ScoreImportResult:
    successful: int = 0
    failed: int = 0
    skipped: int = 0
    updated: int = 0
    total_rows: int = 0
    errors: list[ScoreImportRowError] = field(default_factory=list)
    errors_truncated: bool = False
    # Full error list for artifact export (not truncated)
    all_errors: list[ScoreImportRowError] = field(default_factory=list)
    dry_run: bool = False
    file_checksum: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {
            "successful": self.successful,
            "failed": self.failed,
            "skipped": self.skipped,
            "updated": self.updated,
            "total_rows": self.total_rows,
            "errors": [e.as_dict() for e in self.errors],
            "errors_truncated": self.errors_truncated,
            "dry_run": self.dry_run,
            "file_checksum": self.file_checksum,
        }

    def record_error(self, err: ScoreImportRowError) -> None:
        self.failed += 1
        self.all_errors.append(err)
        if len(self.errors) < MAX_STORED_IMPORT_ERRORS:
            self.errors.append(err)
        else:
            self.errors_truncated = True


@dataclass
class ExamSubjectInfo:
    id: int
    obj_max_score: float | None
    essay_max_score: float | None


@dataclass
class ExistingScoreInfo:
    score_id: int | None
    obj_raw_score: str | None
    essay_raw_score: str | None
    pract_raw_score: str | None


@dataclass
class ImportLookups:
    subject_by_code: dict[str, ExamSubjectInfo]
    reg_by_index: dict[str, int]
    # stripped leading-zero index → unique exam_registration_id (only if unambiguous)
    reg_by_stripped_index: dict[str, int]
    # (exam_registration_id, exam_subject_id) → (subject_registration_id, ExistingScoreInfo)
    sr_lookup: dict[tuple[int, int], tuple[int, ExistingScoreInfo]]


@dataclass
class ReadyApplyRow:
    row_num: int
    subject_registration_id: int
    subject_score_id: int | None
    parsed_score: str
    total_score: float


def file_checksum(file_content: bytes) -> str:
    return hashlib.sha256(file_content).hexdigest()


def new_job_key(*, tracking_id: int | None = None) -> str:
    if tracking_id is not None:
        return f"job-{tracking_id}"
    return f"sync-{uuid.uuid4().hex}"


def _normalize_header(value: Any) -> str:
    import re

    text = str(value or "").strip().lower()
    text = re.sub(r"[\s\-]+", "_", text)
    return text


def normalize_score_import_columns(df: pd.DataFrame) -> pd.DataFrame:
    rename: dict[str, str] = {}
    used_targets: set[str] = set()
    for col in df.columns:
        key = _normalize_header(col)
        target = COLUMN_ALIASES.get(key)
        if target and target not in used_targets:
            rename[col] = target
            used_targets.add(target)
    return df.rename(columns=rename)


def validate_score_import_columns(df: pd.DataFrame) -> None:
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(missing)}. "
            f"Expected: {', '.join(REQUIRED_COLUMNS)}"
        )


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, str) and not value.strip():
        return None
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def _cell_to_raw_str(value: Any) -> str | None:
    """Preserve leading zeros; treat Excel floats carefully."""
    if value is None:
        return None
    if isinstance(value, float):
        if pd.isna(value):
            return None
        if value == int(value):
            return str(int(value))
        return str(value).strip() or None
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def is_skip_score_value(score_text: str | None) -> bool:
    if score_text is None:
        return True
    normalized = score_text.strip().upper()
    if not normalized:
        return True
    return normalized in SKIP_SCORE_TOKENS


def paper_field_name(test_type: PaperTestType) -> str:
    return "obj_raw_score" if test_type == 1 else "essay_raw_score"


def paper_extraction_attr(test_type: PaperTestType) -> str:
    return "obj_extraction_method" if test_type == 1 else "essay_extraction_method"


def paper_max_for_exam_subject(info: ExamSubjectInfo, test_type: PaperTestType) -> float | None:
    if test_type == 1:
        return info.obj_max_score
    return info.essay_max_score


def paper_is_required_for_import(info: ExamSubjectInfo, test_type: PaperTestType) -> bool:
    max_score = paper_max_for_exam_subject(info, test_type)
    return max_score is not None and max_score > 0


def paper_score_is_missing(subject_score: SubjectScore | None, test_type: PaperTestType) -> bool:
    if subject_score is None:
        return True
    raw = getattr(subject_score, paper_field_name(test_type), None)
    if raw is None:
        return True
    if isinstance(raw, str) and not raw.strip():
        return True
    return False


def subject_lookup_keys(subject: Subject) -> set[str]:
    keys: set[str] = set()
    for raw in (subject.original_code, subject.code):
        if raw and str(raw).strip():
            keys.add(str(raw).strip().upper())
    return keys


def _chunked[T](items: list[T], size: int = IN_CHUNK_SIZE):
    for i in range(0, len(items), size):
        yield items[i : i + size]


def _read_xlsx_as_strings(file_content: bytes) -> pd.DataFrame:
    """Stream xlsx with openpyxl read_only; all cells as strings (preserve index zeros)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_content), read_only=True, data_only=True)
    try:
        ws = wb.active
        if ws is None:
            raise SubjectUploadParseError("Excel workbook has no active sheet")
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration as exc:
            raise SubjectUploadParseError("File is empty or contains no data") from exc

        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(header_row)]
        data: list[list[str | None]] = []
        for row in rows_iter:
            if row is None or all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
                continue
            # Pad/truncate to header length
            cells: list[str | None] = []
            for i in range(len(headers)):
                val = row[i] if i < len(row) else None
                cells.append(_cell_to_raw_str(val))
            if all(c is None for c in cells):
                continue
            data.append(cells)
        if not data:
            raise SubjectUploadParseError("File is empty or contains no data")
        return pd.DataFrame(data, columns=headers)
    finally:
        wb.close()


def parse_score_import_file(file_content: bytes, filename: str) -> pd.DataFrame:
    """
    Parse Excel/CSV for score import with string dtypes.

    Legacy .xls is rejected (openpyxl cannot read it reliably).
    Prefer .csv for very large files; .xlsx uses read-only streaming.
    """
    file_lower = (filename or "").lower()
    try:
        if file_lower.endswith(".xls") and not file_lower.endswith(".xlsx"):
            raise SubjectUploadParseError(
                "Legacy .xls is not supported. Save the file as .xlsx or .csv and try again."
            )
        if file_lower.endswith(".csv"):
            df = pd.read_csv(
                io.BytesIO(file_content),
                dtype=str,
                keep_default_na=False,
            )
        elif file_lower.endswith(".xlsx"):
            df = _read_xlsx_as_strings(file_content)
        else:
            raise SubjectUploadParseError(
                f"Unsupported file type. Expected .xlsx or .csv, got {filename}"
            )

        df = df.dropna(how="all")
        if df.empty:
            raise SubjectUploadParseError("File is empty or contains no data")
        return df
    except SubjectUploadParseError:
        raise
    except pd.errors.EmptyDataError as exc:
        raise SubjectUploadParseError("File is empty or contains no data") from exc
    except Exception as exc:
        raise SubjectUploadParseError(f"Failed to parse file: {exc}") from exc


def extract_import_row_tuples(df: pd.DataFrame) -> list[tuple[int, str | None, str | None, str | None]]:
    """Return (excel_row_number, index_number, subject_code, score_text) without iterrows."""
    index_col = df["index_number"].tolist()
    code_col = df["subject_code"].tolist()
    score_col = df["score"].tolist()
    out: list[tuple[int, str | None, str | None, str | None]] = []
    for i in range(len(df)):
        excel_row = i + 2
        out.append(
            (
                excel_row,
                _cell_str(index_col[i]),
                _cell_str(code_col[i]),
                _cell_str(score_col[i]),
            )
        )
    return out


async def load_import_lookups(
    session: AsyncSession,
    *,
    exam_id: int,
    school_id: int | None,
    test_type: PaperTestType,
) -> ImportLookups:
    """Load slim ID maps — no ORM graph of subject registrations."""
    _ = test_type  # reserved for future paper-specific filters

    es_stmt = (
        select(
            ExamSubject.id,
            ExamSubject.obj_max_score,
            ExamSubject.essay_max_score,
            Subject.original_code,
            Subject.code,
        )
        .join(Subject, ExamSubject.subject_id == Subject.id)
        .where(ExamSubject.exam_id == exam_id)
    )
    es_rows = (await session.execute(es_stmt)).all()
    subject_by_code: dict[str, ExamSubjectInfo] = {}
    for es_id, obj_max, essay_max, original_code, code in es_rows:
        info = ExamSubjectInfo(id=es_id, obj_max_score=obj_max, essay_max_score=essay_max)
        for raw in (original_code, code):
            if raw and str(raw).strip():
                subject_by_code[str(raw).strip().upper()] = info

    reg_stmt = (
        select(ExamRegistration.id, ExamRegistration.index_number)
        .join(Candidate, ExamRegistration.candidate_id == Candidate.id)
        .where(ExamRegistration.exam_id == exam_id)
    )
    if school_id is not None:
        reg_stmt = reg_stmt.where(Candidate.school_id == school_id)
    reg_rows = (await session.execute(reg_stmt)).all()

    reg_by_index: dict[str, int] = {}
    stripped_buckets: dict[str, list[int]] = {}
    for reg_id, index_number in reg_rows:
        if not index_number:
            continue
        key = str(index_number).strip()
        reg_by_index[key] = reg_id
        stripped = key.lstrip("0") or "0"
        stripped_buckets.setdefault(stripped, []).append(reg_id)

    reg_by_stripped_index: dict[str, int] = {
        stripped: ids[0] for stripped, ids in stripped_buckets.items() if len(ids) == 1
    }

    reg_ids = list(reg_by_index.values())
    sr_lookup: dict[tuple[int, int], tuple[int, ExistingScoreInfo]] = {}

    for chunk in _chunked(reg_ids):
        sr_stmt = (
            select(
                SubjectRegistration.id,
                SubjectRegistration.exam_registration_id,
                SubjectRegistration.exam_subject_id,
                SubjectScore.id,
                SubjectScore.obj_raw_score,
                SubjectScore.essay_raw_score,
                SubjectScore.pract_raw_score,
            )
            .outerjoin(SubjectScore, SubjectScore.subject_registration_id == SubjectRegistration.id)
            .where(SubjectRegistration.exam_registration_id.in_(chunk))
        )
        for sr_id, er_id, es_id, score_id, obj_raw, essay_raw, pract_raw in (
            await session.execute(sr_stmt)
        ).all():
            sr_lookup[(er_id, es_id)] = (
                sr_id,
                ExistingScoreInfo(
                    score_id=score_id,
                    obj_raw_score=obj_raw,
                    essay_raw_score=essay_raw,
                    pract_raw_score=pract_raw,
                ),
            )

    return ImportLookups(
        subject_by_code=subject_by_code,
        reg_by_index=reg_by_index,
        reg_by_stripped_index=reg_by_stripped_index,
        sr_lookup=sr_lookup,
    )


def resolve_registration_id(lookups: ImportLookups, index_number: str) -> int | None:
    exact = lookups.reg_by_index.get(index_number)
    if exact is not None:
        return exact
    stripped = index_number.lstrip("0") or "0"
    return lookups.reg_by_stripped_index.get(stripped)


def _compute_total_for_apply(
    *,
    test_type: PaperTestType,
    parsed: str,
    existing: ExistingScoreInfo | None,
) -> float:
    if existing is None:
        obj = parsed if test_type == 1 else None
        essay = parsed if test_type == 2 else None
        pract = None
    else:
        obj = parsed if test_type == 1 else existing.obj_raw_score
        essay = parsed if test_type == 2 else existing.essay_raw_score
        pract = existing.pract_raw_score
    return calculate_total_score(obj, essay, pract)


def classify_import_rows(
    rows: list[tuple[int, str | None, str | None, str | None]],
    lookups: ImportLookups,
    *,
    test_type: PaperTestType,
    school_id: int | None = None,
) -> tuple[ScoreImportResult, list[ReadyApplyRow]]:
    """Validate/match rows in pure Python with O(1) maps. No DB I/O."""
    result = ScoreImportResult(total_rows=len(rows))
    ready: list[ReadyApplyRow] = []

    for excel_row, index_number, subject_code_raw, score_text in rows:
        if is_skip_score_value(score_text):
            result.skipped += 1
            continue
        if not index_number:
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    subject_code=subject_code_raw,
                    message="index_number is required",
                )
            )
            continue
        if not subject_code_raw:
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    index_number=index_number,
                    message="subject_code is required",
                )
            )
            continue

        exam_subject = lookups.subject_by_code.get(subject_code_raw.upper())
        if not exam_subject:
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    index_number=index_number,
                    subject_code=subject_code_raw,
                    message="Subject not found for this examination",
                )
            )
            continue
        if not paper_is_required_for_import(exam_subject, test_type):
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    index_number=index_number,
                    subject_code=subject_code_raw,
                    message=f"Paper {test_type} is not required for this subject",
                )
            )
            continue

        exam_reg_id = resolve_registration_id(lookups, index_number)
        if exam_reg_id is None:
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    index_number=index_number,
                    subject_code=subject_code_raw,
                    message="Candidate index number not found for this examination"
                    + (" / school" if school_id is not None else ""),
                )
            )
            continue

        sr_entry = lookups.sr_lookup.get((exam_reg_id, exam_subject.id))
        if sr_entry is None:
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    index_number=index_number,
                    subject_code=subject_code_raw,
                    message="Candidate is not registered for this subject",
                )
            )
            continue

        subject_reg_id, existing = sr_entry
        try:
            parsed = parse_score_value(score_text)
        except ValueError as exc:
            result.record_error(
                ScoreImportRowError(
                    row=excel_row,
                    index_number=index_number,
                    subject_code=subject_code_raw,
                    message=str(exc),
                )
            )
            continue

        max_score = paper_max_for_exam_subject(exam_subject, test_type)
        if max_score is not None and parsed is not None and not is_absent(parsed):
            ok, err = validate_score_range(parsed, max_score)
            if not ok:
                result.record_error(
                    ScoreImportRowError(
                        row=excel_row,
                        index_number=index_number,
                        subject_code=subject_code_raw,
                        message=err or "Score out of range",
                    )
                )
                continue

        if parsed is None:
            result.skipped += 1
            continue

        total = _compute_total_for_apply(
            test_type=test_type, parsed=parsed, existing=existing
        )
        ready.append(
            ReadyApplyRow(
                row_num=excel_row,
                subject_registration_id=subject_reg_id,
                subject_score_id=existing.score_id,
                parsed_score=parsed,
                total_score=total,
            )
        )
        result.successful += 1
        result.updated += 1

    return result, ready


async def _ensure_staging_table(session: AsyncSession) -> None:
    await session.execute(
        text(
            """
            CREATE TABLE IF NOT EXISTS score_import_staging (
                id BIGSERIAL PRIMARY KEY,
                job_key VARCHAR(64) NOT NULL,
                row_num INTEGER NOT NULL,
                subject_registration_id INTEGER,
                subject_score_id INTEGER,
                parsed_score VARCHAR(10),
                total_score DOUBLE PRECISION,
                status VARCHAR(16) NOT NULL DEFAULT 'ready'
            )
            """
        )
    )
    await session.execute(
        text(
            """
            CREATE INDEX IF NOT EXISTS ix_score_import_staging_job_key
            ON score_import_staging (job_key)
            """
        )
    )


async def cleanup_staging(session: AsyncSession, job_key: str) -> None:
    await session.execute(
        text("DELETE FROM score_import_staging WHERE job_key = :job_key"),
        {"job_key": job_key},
    )
    await session.commit()


async def apply_ready_rows_via_staging(
    session: AsyncSession,
    *,
    job_key: str,
    ready: list[ReadyApplyRow],
    test_type: PaperTestType,
    progress_callback: Any | None = None,
    result: ScoreImportResult,
) -> None:
    """Load ready rows into staging and set-based UPDATE/INSERT subject_scores."""
    if not ready:
        if progress_callback:
            await progress_callback(result.total_rows, result)
        return

    await _ensure_staging_table(session)
    await session.execute(
        text("DELETE FROM score_import_staging WHERE job_key = :job_key"),
        {"job_key": job_key},
    )
    await session.commit()

    method_value = METHOD.value
    processed = 0
    total_ready = len(ready)

    for batch in _chunked(ready, IMPORT_APPLY_BATCH_SIZE):
        await session.execute(
            text("DELETE FROM score_import_staging WHERE job_key = :job_key"),
            {"job_key": job_key},
        )

        # Multi-row insert into staging
        values_sql: list[str] = []
        params: dict[str, Any] = {"job_key": job_key}
        for i, row in enumerate(batch):
            values_sql.append(
                f"(:job_key, :row_num_{i}, :sr_id_{i}, :ss_id_{i}, :parsed_{i}, :total_{i}, 'ready')"
            )
            params[f"row_num_{i}"] = row.row_num
            params[f"sr_id_{i}"] = row.subject_registration_id
            params[f"ss_id_{i}"] = row.subject_score_id
            params[f"parsed_{i}"] = row.parsed_score
            params[f"total_{i}"] = row.total_score

        await session.execute(
            text(
                f"""
                INSERT INTO score_import_staging
                    (job_key, row_num, subject_registration_id, subject_score_id,
                     parsed_score, total_score, status)
                VALUES {", ".join(values_sql)}
                """
            ),
            params,
        )

        if test_type == 1:
            await session.execute(
                text(
                    """
                    UPDATE subject_scores AS ss
                    SET
                        obj_raw_score = st.parsed_score,
                        obj_extraction_method = CAST(:method AS dataextractionmethod),
                        total_score = st.total_score,
                        updated_at = :now
                    FROM score_import_staging AS st
                    WHERE st.job_key = :job_key
                      AND st.status = 'ready'
                      AND st.subject_score_id IS NOT NULL
                      AND ss.id = st.subject_score_id
                    """
                ),
                {"job_key": job_key, "method": method_value, "now": datetime.utcnow()},
            )
            await session.execute(
                text(
                    """
                    INSERT INTO subject_scores (
                        subject_registration_id,
                        obj_raw_score,
                        essay_raw_score,
                        pract_raw_score,
                        total_score,
                        obj_extraction_method,
                        created_at,
                        updated_at
                    )
                    SELECT
                        st.subject_registration_id,
                        st.parsed_score,
                        NULL,
                        NULL,
                        st.total_score,
                        CAST(:method AS dataextractionmethod),
                        :now,
                        :now
                    FROM score_import_staging AS st
                    WHERE st.job_key = :job_key
                      AND st.status = 'ready'
                      AND st.subject_score_id IS NULL
                    """
                ),
                {"job_key": job_key, "method": method_value, "now": datetime.utcnow()},
            )
        else:
            await session.execute(
                text(
                    """
                    UPDATE subject_scores AS ss
                    SET
                        essay_raw_score = st.parsed_score,
                        essay_extraction_method = CAST(:method AS dataextractionmethod),
                        total_score = st.total_score,
                        updated_at = :now
                    FROM score_import_staging AS st
                    WHERE st.job_key = :job_key
                      AND st.status = 'ready'
                      AND st.subject_score_id IS NOT NULL
                      AND ss.id = st.subject_score_id
                    """
                ),
                {"job_key": job_key, "method": method_value, "now": datetime.utcnow()},
            )
            await session.execute(
                text(
                    """
                    INSERT INTO subject_scores (
                        subject_registration_id,
                        obj_raw_score,
                        essay_raw_score,
                        pract_raw_score,
                        total_score,
                        essay_extraction_method,
                        created_at,
                        updated_at
                    )
                    SELECT
                        st.subject_registration_id,
                        NULL,
                        st.parsed_score,
                        NULL,
                        st.total_score,
                        CAST(:method AS dataextractionmethod),
                        :now,
                        :now
                    FROM score_import_staging AS st
                    WHERE st.job_key = :job_key
                      AND st.status = 'ready'
                      AND st.subject_score_id IS NULL
                    """
                ),
                {"job_key": job_key, "method": method_value, "now": datetime.utcnow()},
            )

        await session.commit()
        processed += len(batch)

        if progress_callback and (
            processed % IMPORT_PROGRESS_EVERY < IMPORT_APPLY_BATCH_SIZE
            or processed >= total_ready
        ):
            # Map apply progress onto total file rows proportionally for UI
            approx = min(
                result.total_rows,
                int(result.total_rows * (processed / max(total_ready, 1))),
            )
            await progress_callback(approx, result)

    await session.execute(
        text("DELETE FROM score_import_staging WHERE job_key = :job_key"),
        {"job_key": job_key},
    )
    await session.commit()

    if progress_callback:
        await progress_callback(result.total_rows, result)


def build_errors_csv(errors: list[ScoreImportRowError]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(
        buf,
        fieldnames=["row", "index_number", "subject_code", "message"],
        extrasaction="ignore",
    )
    writer.writeheader()
    for err in errors:
        writer.writerow(
            {
                "row": err.row,
                "index_number": err.index_number or "",
                "subject_code": err.subject_code or "",
                "message": err.message,
            }
        )
    return buf.getvalue().encode("utf-8")


async def import_scores_pipeline(
    session: AsyncSession,
    *,
    exam_id: int,
    test_type: PaperTestType,
    file_content: bytes,
    filename: str,
    school_id: int | None = None,
    enforce_row_limit: bool = True,
    progress_callback: Any | None = None,
    dry_run: bool = False,
    job_key: str | None = None,
    preparsed_df: pd.DataFrame | None = None,
) -> ScoreImportResult:
    """
    Full import pipeline: parse → lookups → classify → (optional) staging apply.

    Large files should set enforce_row_limit=False from the async job worker.
    """
    if test_type not in (1, 2):
        raise ValueError("test_type must be 1 (Paper 1) or 2 (Paper 2)")

    exam = (await session.execute(select(Exam).where(Exam.id == exam_id))).scalar_one_or_none()
    if not exam:
        raise ValueError("Examination not found")

    if school_id is not None:
        school = (
            await session.execute(select(School).where(School.id == school_id))
        ).scalar_one_or_none()
        if not school:
            raise ValueError("School not found")

    checksum = file_checksum(file_content)
    key = job_key or new_job_key()

    if preparsed_df is not None:
        df = preparsed_df
    else:
        try:
            df = parse_score_import_file(file_content, filename)
        except SubjectUploadParseError as exc:
            raise ValueError(str(exc)) from exc

    df = normalize_score_import_columns(df)
    validate_score_import_columns(df)

    total_rows = len(df)
    if enforce_row_limit and total_rows > LARGE_IMPORT_ROW_THRESHOLD:
        raise ValueError(
            f"File has {total_rows} rows; maximum for synchronous import is "
            f"{LARGE_IMPORT_ROW_THRESHOLD}. Use async import or filter by school."
        )

    if progress_callback:
        # Signal parse complete
        early = ScoreImportResult(total_rows=total_rows, file_checksum=checksum, dry_run=dry_run)
        await progress_callback(0, early)

    lookups = await load_import_lookups(
        session, exam_id=exam_id, school_id=school_id, test_type=test_type
    )
    rows = extract_import_row_tuples(df)
    result, ready = classify_import_rows(
        rows, lookups, test_type=test_type, school_id=school_id
    )
    result.file_checksum = checksum
    result.dry_run = dry_run

    if progress_callback:
        await progress_callback(total_rows if dry_run else max(total_rows // 2, 1), result)

    if dry_run:
        if progress_callback:
            await progress_callback(total_rows, result)
        return result

    try:
        await apply_ready_rows_via_staging(
            session,
            job_key=key,
            ready=ready,
            test_type=test_type,
            progress_callback=progress_callback,
            result=result,
        )
    finally:
        try:
            await cleanup_staging(session, key)
        except Exception:
            logger.exception("Failed to cleanup score import staging for %s", key)

    return result
