"""Unit tests for score validation report classification and exports."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from app.models import SubjectScore
from app.services.score_validation_report import (
    ReportDetailRow,
    ReportMeta,
    ScoreValidationReportData,
    _school_file_basename,
    _slice_report_for_school,
    build_summary,
    classify_paper,
    classify_score_papers,
    detail_columns_for_status,
    expected_for_invalid,
    generate_validation_report_excel,
    generate_validation_report_pdf,
    group_rows_for_pdf,
    missing_p1_p2_label,
    parse_school_ids,
    parse_statuses,
    parse_test_types,
    primary_report_status,
    resolve_school_ids,
    should_use_report_job,
    status_count_label,
)


def _template_count_label(report_status: str, *, combine_p1_p2: bool = False):
    return lambda n: status_count_label(n, report_status, combine_p1_p2=combine_p1_p2)


def test_status_count_label_normal_and_combine() -> None:
    assert status_count_label(1, "missing") == "1 missing score"
    assert status_count_label(30, "missing") == "30 missing scores"
    assert status_count_label(1, "invalid") == "1 invalid score"
    assert status_count_label(2, "entered") == "2 entered scores"
    assert status_count_label(1, "absent") == "1 absent score"
    assert (
        status_count_label(1, "missing", combine_p1_p2=True)
        == "1 candidate missing papers"
    )
    assert (
        status_count_label(30, "missing", combine_p1_p2=True)
        == "30 candidates missing papers"
    )


def _make_score(**kwargs) -> SubjectScore:
    score = SubjectScore(
        id=1,
        subject_registration_id=1,
        obj_raw_score=None,
        essay_raw_score=None,
        pract_raw_score=None,
        total_score=0.0,
    )
    for key, value in kwargs.items():
        setattr(score, key, value)
    return score


def _make_exam_subject(**kwargs):
    defaults = {
        "obj_max_score": 40.0,
        "essay_max_score": 60.0,
        "pract_max_score": None,
        "pract_pct": None,
        "obj_pct": 40.0,
        "essay_pct": 60.0,
    }
    defaults.update(kwargs)
    return SimpleNamespace(**defaults)


def _detail_row(**overrides) -> ReportDetailRow:
    base = dict(
        school_id=1,
        school_code="SCH01",
        school_name="Sample School",
        subject_id=1,
        subject_code="MATH",
        subject_name="Mathematics",
        subject_type="CORE",
        candidate_id=1,
        index_number="0011223344",
        candidate_name="Test Candidate",
        test_type=1,
        paper_label="Paper 1 (Objectives)",
        paper_short="P1",
        raw_score=None,
        max_score=40.0,
        status="missing",
        message="Objectives score is missing. Maximum score is 40",
        extraction_method=None,
        expected=None,
    )
    base.update(overrides)
    return ReportDetailRow(**base)


def test_classify_missing_and_entered() -> None:
    score = _make_score(obj_raw_score=None, essay_raw_score="45")
    exam_subject = _make_exam_subject()
    status, message, expected = classify_paper(score, exam_subject, 1)
    assert status == "missing"
    assert message and "missing" in message.lower()
    assert expected is None

    status2, message2, expected2 = classify_paper(score, exam_subject, 2)
    assert status2 == "entered"
    assert message2 is None
    assert expected2 is None


def test_classify_invalid_and_absent() -> None:
    score = _make_score(obj_raw_score="99", essay_raw_score="A")
    exam_subject = _make_exam_subject(obj_max_score=40.0, essay_max_score=60.0)
    status, message, expected = classify_paper(score, exam_subject, 1)
    assert status == "invalid"
    assert message
    assert expected == "0–40"

    status2, _, expected2 = classify_paper(score, exam_subject, 2)
    assert status2 == "absent"
    assert expected2 is None


def test_expected_for_invalid_whole_number_and_range() -> None:
    assert expected_for_invalid(message="Must be a whole number", max_score=40.0) == "Whole number"
    assert expected_for_invalid(message="Score exceeds maximum", max_score=60.0) == "0–60"
    assert expected_for_invalid(message="bad", max_score=None) == "Valid score"


def test_classify_score_papers_filters_test_types() -> None:
    score = _make_score(obj_raw_score="10", essay_raw_score=None, pract_raw_score="A")
    exam_subject = _make_exam_subject(pract_max_score=20.0)
    papers = classify_score_papers(score, exam_subject, test_types={1, 3})
    assert [p[0] for p in papers] == [1, 3]
    assert papers[0][1] == "entered"
    assert papers[1][1] == "absent"
    assert papers[0][5] is None  # expected


def test_classify_skips_non_required_papers() -> None:
    score = _make_score()
    exam_subject = _make_exam_subject(
        obj_max_score=40.0,
        essay_max_score=None,
        pract_max_score=None,
        pract_pct=None,
    )
    papers = classify_score_papers(score, exam_subject)
    assert len(papers) == 1
    assert papers[0][0] == 1


def test_missing_p1_p2_label_both() -> None:
    score = _make_score(obj_raw_score=None, essay_raw_score=None)
    exam_subject = _make_exam_subject()
    assert missing_p1_p2_label(score, exam_subject) == "P1/P2"


def test_missing_p1_p2_label_only_p1() -> None:
    score = _make_score(obj_raw_score=None, essay_raw_score="40")
    exam_subject = _make_exam_subject()
    assert missing_p1_p2_label(score, exam_subject) == "P1"


def test_missing_p1_p2_label_only_p2() -> None:
    score = _make_score(obj_raw_score="10", essay_raw_score=None)
    exam_subject = _make_exam_subject()
    assert missing_p1_p2_label(score, exam_subject) == "P2"


def test_missing_p1_p2_label_neither() -> None:
    score = _make_score(obj_raw_score="10", essay_raw_score="40")
    exam_subject = _make_exam_subject()
    assert missing_p1_p2_label(score, exam_subject) is None


def test_missing_p1_p2_label_only_p1_required() -> None:
    score = _make_score(obj_raw_score=None, essay_raw_score=None)
    exam_subject = _make_exam_subject(essay_max_score=None)
    assert missing_p1_p2_label(score, exam_subject) == "P1"


def test_parse_statuses_single_required() -> None:
    assert parse_statuses(None) == ["missing"]
    assert parse_statuses("") == ["missing"]
    assert parse_statuses("invalid") == ["invalid"]
    with pytest.raises(ValueError, match="exactly one"):
        parse_statuses("all")
    with pytest.raises(ValueError, match="exactly one"):
        parse_statuses("entered,absent")
    with pytest.raises(ValueError):
        parse_statuses("foo")


def test_parse_test_types() -> None:
    assert parse_test_types(None) is None
    assert parse_test_types("1,2") == [1, 2]
    with pytest.raises(ValueError):
        parse_test_types("9")


def test_detail_columns_by_status() -> None:
    missing = [c.header for c in detail_columns_for_status("missing")]
    assert missing == ["Index number", "Candidate name"]

    combined = [c.header for c in detail_columns_for_status("missing", combine_p1_p2=True)]
    assert combined == ["Index number", "Candidate name", "Missing papers"]

    invalid = [c.header for c in detail_columns_for_status("invalid")]
    assert invalid == ["Index number", "Candidate name", "Value", "Expected"]

    absent = [c.header for c in detail_columns_for_status("absent")]
    assert absent == ["Index number", "Candidate name"]

    entered = [c.header for c in detail_columns_for_status("entered")]
    assert entered == ["Index number", "Candidate name", "Score"]


def test_build_summary_and_sort_groups() -> None:
    rows = [
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            index_number="2",
            candidate_name="Bob",
            status="missing",
            message="missing",
        ),
        _detail_row(
            school_id=1,
            school_code="A01",
            school_name="Alpha",
            candidate_id=2,
            index_number="1",
            candidate_name="Ann",
            raw_score="12",
            status="entered",
            message=None,
        ),
        _detail_row(
            school_id=1,
            school_code="A01",
            school_name="Alpha",
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            candidate_id=2,
            index_number="1",
            candidate_name="Ann",
            test_type=2,
            paper_label="Paper 2 (Essay)",
            paper_short="P2",
            raw_score="999",
            max_score=60.0,
            status="invalid",
            message="too high",
            expected="0–60",
        ),
    ]
    rows.sort(key=lambda r: (r.school_code, r.test_type, r.subject_code, r.index_number))
    summary = build_summary(rows)
    assert summary.total == 3
    assert summary.entered == 1
    assert summary.missing == 1
    assert summary.invalid == 1
    assert summary.by_school[0]["school_code"] == "A01"
    sections = group_rows_for_pdf(rows)
    assert sections[0]["school_code"] == "A01"
    # Papers are separate top-level groups under the school
    assert [p["test_type"] for p in sections[0]["papers"]] == [1, 2]
    assert sections[0]["papers"][0]["subjects"][0]["subject_code"] in {"MATH", "ENG"}


def test_papers_not_mixed_in_sections() -> None:
    rows = [
        _detail_row(test_type=1, paper_short="P1", paper_label="Paper 1 (Objectives)", status="missing"),
        _detail_row(
            test_type=2,
            paper_short="P2",
            paper_label="Paper 2 (Essay)",
            max_score=60.0,
            status="missing",
            candidate_id=2,
            index_number="2",
        ),
        _detail_row(
            test_type=1,
            paper_short="P1",
            paper_label="Paper 1 (Objectives)",
            status="missing",
            candidate_id=3,
            index_number="3",
        ),
    ]
    rows.sort(key=lambda r: (r.school_code, r.test_type, r.subject_code, r.index_number))
    sections = group_rows_for_pdf(rows)
    papers = sections[0]["papers"]
    assert len(papers) == 2
    assert papers[0]["test_type"] == 1
    assert papers[1]["test_type"] == 2
    assert all(r.test_type == 1 for r in papers[0]["subjects"][0]["rows"])
    assert all(r.test_type == 2 for r in papers[1]["subjects"][0]["rows"])
    assert papers[0]["subjects"][0]["max_score"] == "40"
    assert papers[1]["subjects"][0]["max_score"] == "60"


def test_parse_and_resolve_school_ids() -> None:
    assert parse_school_ids(None) is None
    assert parse_school_ids("") is None
    assert parse_school_ids([]) is None
    assert parse_school_ids("1,2,3") == [1, 2, 3]
    assert parse_school_ids([4, 5]) == [4, 5]
    assert resolve_school_ids(school_ids="10,20") == [10, 20]
    assert resolve_school_ids(school_ids=None, school_id=7) == [7]
    assert resolve_school_ids(school_ids="1,2", school_id=7) == [1, 2]
    assert resolve_school_ids() is None


def test_should_use_report_job() -> None:
    assert should_use_report_job(school_id=None, report_format="pdf") is True
    assert should_use_report_job(school_ids=None, report_format="xlsx") is True
    assert should_use_report_job(school_ids=[1, 2], report_format="xlsx") is True
    assert should_use_report_job(school_ids=[], report_format="xlsx") is True
    assert should_use_report_job(school_id=1, report_format="xlsx", estimated_rows=100) is False
    assert should_use_report_job(school_ids=[1], report_format="xlsx", estimated_rows=100) is False
    assert should_use_report_job(school_id=1, report_format="xlsx", estimated_rows=6000) is True
    assert should_use_report_job(school_id=1, report_format="xlsx", school_count=3) is True
    # Sync PDF path must reject before render when over the PDF row threshold.
    assert should_use_report_job(school_id=1, report_format="pdf", estimated_rows=1500) is False
    assert should_use_report_job(school_id=1, report_format="pdf", estimated_rows=1501) is True
    assert should_use_report_job(school_id=1, report_format="xlsx", estimated_rows=1501) is False
    # school_ids takes precedence over singular school_id
    assert should_use_report_job(
        school_id=1, school_ids=[1, 2], report_format="xlsx", estimated_rows=10
    ) is True


def _sample_report_data(
    *,
    status: str = "missing",
    rows: list[ReportDetailRow] | None = None,
) -> ScoreValidationReportData:
    if rows is None:
        rows = [
            _detail_row(
                status=status,
                message="Objectives score is missing. Maximum score is 40"
                if status == "missing"
                else None,
                expected="0–40" if status == "invalid" else None,
                raw_score="99" if status == "invalid" else ("12" if status == "entered" else None),
            )
        ]
    meta = ReportMeta(
        exam_id=1,
        exam_label="2026 MAY/JUNE CERT2",
        exam_year=2026,
        exam_series="MAY/JUNE",
        exam_type="CERT2",
        school_id=rows[0].school_id if len({r.school_id for r in rows}) == 1 else None,
        school_label=f"{rows[0].school_code} — {rows[0].school_name}"
        if len({r.school_id for r in rows}) == 1
        else None,
        school_code=rows[0].school_code if len({r.school_id for r in rows}) == 1 else None,
        subject_type="CORE",
        subject_ids=None,
        subject_labels=[],
        test_types=[1],
        statuses=[status],
        generated_at="2026-08-30 12:00 UTC",
        row_count=len(rows),
    )
    return ScoreValidationReportData(meta=meta, summary=build_summary(rows), rows=rows)


def test_excel_export_adaptive_columns_and_subject_sheets() -> None:
    rows = [
        _detail_row(status="invalid", raw_score="99", expected="0–40", message="too high"),
        _detail_row(
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            status="invalid",
            raw_score="abc",
            expected="Whole number",
            message="Must be a whole number",
            test_type=2,
            paper_label="Paper 2 (Essay)",
            paper_short="P2",
            max_score=60.0,
        ),
    ]
    data = _sample_report_data(status="invalid", rows=rows)
    excel_bytes = generate_validation_report_excel(data)
    assert excel_bytes[:2] == b"PK"

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Summary" not in wb.sheetnames
    math_ws = wb["P1_MATH"]
    # Row 1 title, 2 centre, 3 summary, 5 headers, 6+ data
    assert math_ws.cell(2, 1).value == "Centre"
    assert "SCH01" in (math_ws.cell(2, 2).value or "")
    assert math_ws.cell(3, 1).value == "Summary"
    assert "1 invalid score" in (math_ws.cell(3, 2).value or "")
    assert "row" not in (math_ws.cell(3, 2).value or "").lower()
    headers = [math_ws.cell(5, c).value for c in range(1, 6)]
    assert headers == ["#", "Index number", "Candidate name", "Value", "Expected"]
    assert math_ws.cell(6, 1).value == 1
    assert math_ws.cell(6, 4).value == "99"
    assert math_ws.cell(6, 5).value == "0–40"
    assert "MATH" in (math_ws.cell(1, 1).value or "")

    assert "P2_ENG" in wb.sheetnames
    assert "MATH" not in wb.sheetnames  # no mixed paper-free subject sheet

    missing_data = _sample_report_data(status="missing")
    missing_wb = openpyxl.load_workbook(io.BytesIO(generate_validation_report_excel(missing_data)))
    assert "Summary" not in missing_wb.sheetnames
    missing_headers = [missing_wb["P1_MATH"].cell(5, c).value for c in range(1, 5)]
    assert missing_headers == ["#", "Index number", "Candidate name", None]


def test_excel_multi_school_separate_subject_sheets() -> None:
    """Merged multi-school workbook: one sheet per school+subject, no global Summary."""
    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
        _detail_row(
            school_id=1,
            school_code="A01",
            school_name="Alpha",
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            candidate_id=3,
            index_number="3",
            candidate_name="Cara",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    wb = openpyxl.load_workbook(io.BytesIO(generate_validation_report_excel(data)))
    assert "Summary" not in wb.sheetnames
    assert "A01_P1_MATH" in wb.sheetnames
    assert "B01_P1_MATH" in wb.sheetnames
    assert "A01_P1_ENG" in wb.sheetnames
    # Schools must not share a paper+subject sheet
    assert "P1_MATH" not in wb.sheetnames

    a_math = wb["A01_P1_MATH"]
    assert "Alpha" in (a_math.cell(2, 2).value or "")
    assert "1 missing score" in (a_math.cell(3, 2).value or "")
    assert a_math.cell(6, 1).value == 1

    b_math = wb["B01_P1_MATH"]
    assert "Beta" in (b_math.cell(2, 2).value or "")
    assert "1 missing score" in (b_math.cell(3, 2).value or "")


def test_pdf_export_subject_page_breaks(monkeypatch: pytest.MonkeyPatch) -> None:
    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )
    assert (templates / "score_validation_report" / "main.html").is_file()

    rows = [
        _detail_row(status="missing"),
        _detail_row(
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            status="missing",
            message="Essay score is missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    pdf_bytes = generate_validation_report_pdf(data)
    assert pdf_bytes[:4] == b"%PDF"
    assert len(pdf_bytes) > 100

    # Rendered HTML path uses page-break between subjects — assert via template render.
    from jinja2 import Environment, FileSystemLoader

    from app.services.score_validation_report import detail_columns_for_status, primary_report_status

    env = Environment(loader=FileSystemLoader(str(templates)))
    template = env.get_template("score_validation_report/main.html")
    report_status = primary_report_status(data.meta.statuses)
    html = template.render(
        meta=data.meta,
        summary=data.summary,
        sections=group_rows_for_pdf(data.rows),
        report_status=report_status,
        status_count=getattr(data.summary, report_status),
        columns=detail_columns_for_status(report_status),
        logo_src="score_sheets/logo-crest-only.png",
        status_labels={
            "entered": "Entered",
            "missing": "Missing",
            "invalid": "Invalid",
            "absent": "Absent",
        },
        status_count_label=_template_count_label(report_status),
    )
    assert "page-break" in html
    assert "MATH" in html and "ENG" in html
    assert "Score correction worksheet" in html
    assert "summary-strip" not in html  # no aggregate cover summary
    assert html.count('class="block-summary"') == 2  # one per subject
    assert "1 missing score" in html
    assert "row(s)" not in html
    assert 'content: "Page " counter(page) " / " counter(pages)' in (
        templates / "score_validation_report" / "main.html"
    ).read_text()
    assert "width: 44px" in (templates / "score_validation_report" / "main.html").read_text()
    assert "col-score" in html
    assert ">Score</th>" in html or "Correct score" in html
    assert "Maximum mark" in html
    assert "Initials" not in html
    assert "Validation sign-off" not in html
    assert "Issue message" not in html
    assert "Extraction method" not in html
    assert "A4 portrait" in (templates / "score_validation_report" / "main.html").read_text()


def test_pdf_no_aggregate_cover_multi_school(monkeypatch: pytest.MonkeyPatch) -> None:
    """Each school+subject gets its own summary; no single cover listing all blocks."""
    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )
    from jinja2 import Environment, FileSystemLoader

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    env = Environment(loader=FileSystemLoader(str(templates)))
    template = env.get_template("score_validation_report/main.html")
    html = template.render(
        meta=data.meta,
        summary=data.summary,
        sections=group_rows_for_pdf(data.rows),
        report_status="missing",
        status_count=2,
        columns=detail_columns_for_status("missing"),
        logo_src="score_sheets/logo-crest-only.png",
        status_labels={
            "entered": "Entered",
            "missing": "Missing",
            "invalid": "Invalid",
            "absent": "Absent",
        },
        status_count_label=_template_count_label("missing"),
    )
    assert "summary-strip" not in html
    assert "Examinations · Score entry" not in html  # old cover kicker
    assert html.count('class="block-summary"') == 2
    assert "A01 — Alpha" in html
    assert "B01 — Beta" in html
    assert "MATH" in html and "ENG" in html


def test_pdf_invalid_has_correct_score_column(monkeypatch: pytest.MonkeyPatch) -> None:
    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )
    from jinja2 import Environment, FileSystemLoader

    rows = [
        _detail_row(status="invalid", raw_score="99", expected="0–40", message="too high"),
    ]
    data = _sample_report_data(status="invalid", rows=rows)
    env = Environment(loader=FileSystemLoader(str(templates)))
    template = env.get_template("score_validation_report/main.html")
    html = template.render(
        meta=data.meta,
        summary=data.summary,
        sections=group_rows_for_pdf(data.rows),
        report_status="invalid",
        status_count=1,
        columns=detail_columns_for_status("invalid"),
        logo_src="score_sheets/logo-crest-only.png",
        status_labels={
            "entered": "Entered",
            "missing": "Missing",
            "invalid": "Invalid",
            "absent": "Absent",
        },
        status_count_label=_template_count_label("invalid"),
    )
    assert "Correct score" in html
    assert ">Expected</th>" not in html
    assert "bad-value" in html
    assert "1 invalid score" in html
    assert "row(s)" not in html
    pdf_bytes = generate_validation_report_pdf(data)
    assert pdf_bytes[:4] == b"%PDF"


def test_pdf_combine_p1_p2_has_dual_score_columns(monkeypatch: pytest.MonkeyPatch) -> None:
    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )
    from jinja2 import Environment, FileSystemLoader

    rows = [
        _detail_row(
            test_type=0,
            paper_label="Paper 1 or 2 (combined)",
            paper_short="P1/P2",
            status="missing",
            missing_papers="P1/P2",
            message="Missing score(s): P1/P2",
            raw_score=None,
            max_score=None,
        ),
        _detail_row(
            candidate_id=2,
            index_number="002",
            candidate_name="Only P1 Missing",
            test_type=0,
            paper_label="Paper 1 or 2 (combined)",
            paper_short="P1",
            status="missing",
            missing_papers="P1",
            message="Missing score(s): P1",
            raw_score=None,
            max_score=None,
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    data.meta.combine_p1_p2 = True
    data.meta.test_types = [1, 2]

    env = Environment(loader=FileSystemLoader(str(templates)))
    template = env.get_template("score_validation_report/main.html")
    html = template.render(
        meta=data.meta,
        summary=data.summary,
        sections=group_rows_for_pdf(data.rows),
        report_status="missing",
        status_count=2,
        columns=detail_columns_for_status("missing", combine_p1_p2=True),
        combine_p1_p2=True,
        logo_src="score_sheets/logo-crest-only.png",
        status_labels={
            "entered": "Entered",
            "missing": "Missing",
            "invalid": "Invalid",
            "absent": "Absent",
        },
        status_count_label=_template_count_label("missing", combine_p1_p2=True),
    )
    assert "Missing papers" in html
    assert ">P1</th>" in html
    assert ">P2</th>" in html
    assert html.count("col-score-dual write-box") >= 2
    assert "P1 and/or <strong>P2</strong>" in html or "P1</strong> and/or <strong>P2" in html
    assert "Write-in: P1 and P2 score columns" in html
    assert "block-summary" in html
    assert "2 candidates missing papers" in html
    assert "missing scores" not in html
    assert "row(s)" not in html

    excel = openpyxl.load_workbook(io.BytesIO(generate_validation_report_excel(data)))
    # Combined sheet names use P1P2
    combined_sheet = next(n for n in excel.sheetnames if "P1P2" in n or "MATH" in n)
    summary = excel[combined_sheet].cell(3, 2).value or ""
    assert "candidates missing papers" in summary
    assert "missing scores" not in summary
    assert "row" not in summary.lower()

    pdf_bytes = generate_validation_report_pdf(data)
    assert isinstance(pdf_bytes, (bytes, bytearray))
    assert len(pdf_bytes) > 100
    assert pdf_bytes[:4] == b"%PDF"


def test_slice_and_multi_school_zip_structure() -> None:
    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    assert primary_report_status(data.meta.statuses) == "missing"

    school_a = _slice_report_for_school(data, 1)
    assert school_a.meta.school_id == 1
    assert school_a.meta.school_code == "A01"
    assert len(school_a.rows) == 1

    school_b = _slice_report_for_school(data, 2)
    name_a = _school_file_basename(school_a, "xlsx")
    name_b = _school_file_basename(school_b, "xlsx")
    assert name_a.startswith("A01_Alpha_") and name_a.endswith(".xlsx")
    assert name_b.startswith("B01_Beta_") and name_b.endswith(".xlsx")
    assert name_a != name_b

    # Package like generate_score_validation_report_bytes multi-school path
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for sid in (1, 2):
            school_data = _slice_report_for_school(data, sid)
            entry = _school_file_basename(school_data, "xlsx")
            zf.writestr(entry, generate_validation_report_excel(school_data))

    zip_buf.seek(0)
    with zipfile.ZipFile(zip_buf, "r") as zf:
        names = zf.namelist()
    assert len(names) == 2
    assert any("A01" in n for n in names)
    assert any("B01" in n for n in names)


def test_merged_multi_school_pdf_is_single_document() -> None:
    """Merged packaging renders one PDF covering all schools (no zip)."""
    from PyPDF2 import PdfReader

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    pdf_bytes = generate_validation_report_pdf(data)
    assert pdf_bytes[:4] == b"%PDF"
    assert len(pdf_bytes) > 100
    # Not a zip
    assert pdf_bytes[:2] != b"PK"
    # One WeasyPrint doc per school, then merge → one page per school here
    assert len(PdfReader(io.BytesIO(pdf_bytes)).pages) == 2


def test_pdf_one_render_per_school_multi_subject(monkeypatch: pytest.MonkeyPatch) -> None:
    """Multiple subjects in one school share a single WeasyPrint document."""
    from PyPDF2 import PdfReader

    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )

    from app.services import score_validation_report as svr

    render_calls: list[int] = []
    real_render = svr._render_validation_report_pdf_html

    def tracking_render(data):  # type: ignore[no-untyped-def]
        render_calls.append(len(data.rows))
        return real_render(data)

    monkeypatch.setattr(svr, "_render_validation_report_pdf_html", tracking_render)

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=1,
            school_code="A01",
            school_name="Alpha",
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
        _detail_row(
            school_id=1,
            school_code="A01",
            school_name="Alpha",
            subject_id=3,
            subject_code="SCI",
            subject_name="Science",
            candidate_id=3,
            index_number="3",
            candidate_name="Cara",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    pdf_bytes = generate_validation_report_pdf(data)
    assert pdf_bytes[:4] == b"%PDF"
    assert len(render_calls) == 1
    assert render_calls[0] == 3
    # Subjects still page-break inside the school document
    assert len(PdfReader(io.BytesIO(pdf_bytes)).pages) == 3


def test_pdf_page_numbering_resets_per_school(monkeypatch: pytest.MonkeyPatch) -> None:
    """Two schools → two WeasyPrint docs merged; page counters reset at school boundary."""
    from PyPDF2 import PdfReader

    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )

    from app.services.score_validation_report import (
        _iter_school_row_groups,
        _render_validation_report_pdf_html,
        _slice_report_for_rows,
    )

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=1,
            school_code="A01",
            school_name="Alpha",
            subject_id=2,
            subject_code="ENG",
            subject_name="English",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=3,
            index_number="3",
            candidate_name="Cara",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    groups = _iter_school_row_groups(data.rows)
    assert len(groups) == 2
    assert len(groups[0]) == 2
    assert len(groups[1]) == 1

    part_a = _render_validation_report_pdf_html(_slice_report_for_rows(data, groups[0]))
    part_b = _render_validation_report_pdf_html(_slice_report_for_rows(data, groups[1]))
    assert len(PdfReader(io.BytesIO(part_a)).pages) == 2
    assert len(PdfReader(io.BytesIO(part_b)).pages) == 1

    merged = generate_validation_report_pdf(data)
    assert len(PdfReader(io.BytesIO(merged)).pages) == 3


def test_pdf_worker_count_bounded() -> None:
    from app.services.score_validation_report import PDF_RENDER_MAX_WORKERS, _pdf_worker_count

    assert _pdf_worker_count(0) == 1
    assert _pdf_worker_count(1) == 1
    assert _pdf_worker_count(100) <= PDF_RENDER_MAX_WORKERS
    assert PDF_RENDER_MAX_WORKERS == 2
    assert _pdf_worker_count(100) >= 1


def test_jinja_template_cached_across_renders(monkeypatch: pytest.MonkeyPatch) -> None:
    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )
    from app.services.score_validation_report import (
        _get_validation_report_template,
        _jinja_cache,
        _logo_data_uri,
    )

    _jinja_cache["path"] = None
    _jinja_cache["template"] = None
    _jinja_cache["logo_data_uri"] = None
    t1, d1 = _get_validation_report_template()
    t2, d2 = _get_validation_report_template()
    assert t1 is t2
    assert d1 == d2
    assert d1 == templates.resolve()
    logo1 = _logo_data_uri(d1)
    logo2 = _logo_data_uri(d1)
    assert logo1 == logo2
    assert logo1.startswith("data:image/png;base64,")


def test_multi_school_pdf_uses_thread_pool(monkeypatch: pytest.MonkeyPatch) -> None:
    """Multi-school sync PDF merge uses ThreadPoolExecutor (not ProcessPool)."""
    from concurrent.futures import ProcessPoolExecutor, ThreadPoolExecutor

    from PyPDF2 import PdfReader

    templates = Path(__file__).resolve().parents[1] / "templates"
    monkeypatch.setattr(
        "app.services.score_validation_report.settings.templates_path",
        str(templates),
    )

    from app.services import score_validation_report as svr

    thread_submits: list[int] = []
    real_submit = ThreadPoolExecutor.submit

    def tracking_submit(self, fn, *args, **kwargs):  # type: ignore[no-untyped-def]
        thread_submits.append(1)
        return real_submit(self, fn, *args, **kwargs)

    def boom_process_pool(*_args, **_kwargs):  # type: ignore[no-untyped-def]
        raise AssertionError("ProcessPoolExecutor must not be used for PDF render")

    monkeypatch.setattr(ThreadPoolExecutor, "submit", tracking_submit)
    monkeypatch.setattr(svr, "ProcessPoolExecutor", boom_process_pool, raising=False)
    monkeypatch.setattr(ProcessPoolExecutor, "__init__", boom_process_pool)
    monkeypatch.setattr(svr, "_pdf_worker_count", lambda n: min(2, n) if n > 1 else 1)

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
        _detail_row(
            school_id=3,
            school_code="C01",
            school_name="Gamma",
            candidate_id=3,
            index_number="3",
            candidate_name="Cara",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    pdf_bytes = generate_validation_report_pdf(data)
    assert pdf_bytes[:4] == b"%PDF"
    assert len(thread_submits) == 3
    assert len(PdfReader(io.BytesIO(pdf_bytes)).pages) == 3


@pytest.mark.asyncio
async def test_merged_pdf_emits_intermediate_progress(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Merged PDF must tick schools_done during render, not only at the end."""
    from app.services import score_validation_report as svr

    monkeypatch.setattr(svr, "_pdf_worker_count", lambda _n: 1)

    async def fake_filename(*_args, **_kwargs):
        return "merged_score_validation_MISSING.pdf"

    def fake_render(school_data, report_format):  # type: ignore[no-untyped-def]
        assert report_format == "pdf"
        # Minimal valid-ish PDF header for merge
        return b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF\n"

    monkeypatch.setattr(svr, "generate_report_filename", fake_filename)
    monkeypatch.setattr(svr, "_render_report_file", fake_render)
    monkeypatch.setattr(
        svr,
        "_merge_pdf_parts",
        lambda parts: b"%PDF-merged-" + str(len(parts)).encode() + b"\n",
    )

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
        _detail_row(
            school_id=3,
            school_code="C01",
            school_name="Gamma",
            candidate_id=3,
            index_number="3",
            candidate_name="Cara",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    progress: dict = {}
    seen_done: list[int] = []

    async def on_progress() -> None:
        seen_done.append(int(progress.get("schools_done") or 0))

    from unittest.mock import MagicMock

    file_bytes, filename = await svr.render_score_validation_report_from_data(
        MagicMock(),
        data,
        exam_id=1,
        report_format="pdf",
        packaging="merged",
        progress=progress,
        on_progress=on_progress,
    )
    assert filename.endswith(".pdf")
    assert file_bytes.startswith(b"%PDF-merged-3")
    assert 1 in seen_done
    assert 2 in seen_done
    assert 3 in seen_done
    assert progress.get("schools_done") == 3


@pytest.mark.asyncio
async def test_merged_pdf_coerced_to_zip_over_school_limit(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Merged PDF with > MERGED_PDF_SCHOOL_LIMIT schools becomes a zip."""
    from app.services import score_validation_report as svr

    monkeypatch.setattr(svr, "MERGED_PDF_SCHOOL_LIMIT", 2)
    monkeypatch.setattr(svr, "_pdf_worker_count", lambda _n: 1)

    async def fake_filename(*_args, **_kwargs):
        return "all_schools_score_validation_MISSING.pdf"

    monkeypatch.setattr(svr, "generate_report_filename", fake_filename)
    monkeypatch.setattr(
        svr,
        "_render_report_file",
        lambda _data, _fmt: b"%PDF-school\n",
    )

    rows = [
        _detail_row(school_id=1, school_code="A01", school_name="Alpha", status="missing"),
        _detail_row(
            school_id=2,
            school_code="B01",
            school_name="Beta",
            candidate_id=2,
            index_number="2",
            candidate_name="Bob",
            status="missing",
        ),
        _detail_row(
            school_id=3,
            school_code="C01",
            school_name="Gamma",
            candidate_id=3,
            index_number="3",
            candidate_name="Cara",
            status="missing",
        ),
    ]
    data = _sample_report_data(status="missing", rows=rows)
    progress: dict = {}
    messages: list[str] = []

    async def on_progress() -> None:
        messages.append(str(progress.get("message") or ""))

    from unittest.mock import MagicMock

    file_bytes, filename = await svr.render_score_validation_report_from_data(
        MagicMock(),
        data,
        exam_id=1,
        report_format="pdf",
        packaging="merged",
        progress=progress,
        on_progress=on_progress,
    )
    assert filename.endswith(".zip")
    assert file_bytes[:2] == b"PK"
    assert progress.get("packaging_coerced") == "zip"
    assert any("zip" in m.lower() for m in messages)


def test_classify_score_papers_validates_once(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[int] = []
    from app.services import score_validation_report as svr

    real = svr.validate_subject_score

    def tracking_validate(score, exam_subject):  # type: ignore[no-untyped-def]
        calls.append(1)
        return real(score, exam_subject)

    monkeypatch.setattr(svr, "validate_subject_score", tracking_validate)
    score = _make_score()
    exam_subject = _make_exam_subject()
    papers = classify_score_papers(score, exam_subject)
    assert len(papers) >= 2
    assert len(calls) == 1


def test_early_size_rejection_skips_render(monkeypatch: pytest.MonkeyPatch) -> None:
    """Sync download must call should_use_report_job before PDF/Excel render."""
    from unittest.mock import MagicMock

    import asyncio

    from app.routers import scores as scores_router

    rows = [_detail_row(status="missing") for _ in range(3)]
    data = _sample_report_data(status="missing", rows=rows)
    # Force meta.row_count over PDF sync threshold via monkeypatched should_use
    render_called = {"value": False}

    async def fake_build(*_args, **_kwargs):
        return data

    async def fake_render(*_args, **_kwargs):
        render_called["value"] = True
        return b"%PDF-fake", "report.pdf"

    monkeypatch.setattr(scores_router, "build_score_validation_report", fake_build)
    monkeypatch.setattr(scores_router, "render_score_validation_report_from_data", fake_render)
    monkeypatch.setattr(
        scores_router,
        "should_use_report_job",
        lambda **_kwargs: True,
    )

    session = MagicMock()
    user = MagicMock()

    async def run() -> None:
        with pytest.raises(scores_router.HTTPException) as exc_info:
            await scores_router.download_score_validation_report(
                session=session,
                _user=user,
                exam_id=1,
                school_id=1,
                school_ids=None,
                subject_type=None,
                subject_ids=None,
                test_types=None,
                statuses=None,
                combine_p1_p2=False,
                format="pdf",
                packaging="zip",
            )
        assert exc_info.value.status_code == 413
        assert render_called["value"] is False

    asyncio.run(run())
